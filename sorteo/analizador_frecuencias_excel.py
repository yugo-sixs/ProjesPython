from collections import Counter
from datetime import datetime
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COLUMNA_CONCURSO = 1  # A
COLUMNAS_A_ANALIZAR = range(3, 10)  # C hasta I


def obtener_carpeta_salida():
    """Devuelve la carpeta sorteo donde se guardan los reportes."""
    carpeta_programa = Path(__file__).resolve().parent
    if carpeta_programa.name.lower() == "sorteo":
        return carpeta_programa
    return carpeta_programa.parent / "sorteo"


def normalizar_numero(valor):
    """Convierte celdas numericas a una forma estable para el conteo."""
    if valor is None:
        return None

    if isinstance(valor, bool):
        return None

    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and valor.is_integer():
            return int(valor)
        return valor

    if isinstance(valor, str):
        texto = valor.strip()
        if not texto:
            return None

        texto = texto.replace(",", ".")
        try:
            numero = float(texto)
        except ValueError:
            return None

        if numero.is_integer():
            return int(numero)
        return numero

    return None


def ajustar_columnas(hoja):
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = 12
        for celda in columna:
            if celda.value is not None:
                ancho = max(ancho, len(str(celda.value)) + 2)
        hoja.column_dimensions[letra].width = min(ancho, 40)


def aplicar_estilo_encabezado(hoja):
    relleno = PatternFill("solid", fgColor="1F4E78")
    fuente = Font(color="FFFFFF", bold=True)
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = Alignment(horizontal="center")


def obtener_lugares_frecuencia(contador, cantidad_lugares=3):
    lugares = []
    frecuencias = sorted(set(contador.values()), reverse=True)

    for frecuencia in frecuencias[:cantidad_lugares]:
        numeros_lugar = [
            numero
            for numero, veces in sorted(contador.items(), key=lambda item: item[0])
            if veces == frecuencia
        ]
        numeros = [
            str(numero) for numero in numeros_lugar
        ]
        lugares.append((", ".join(numeros), frecuencia, set(numeros_lugar)))

    while len(lugares) < cantidad_lugares:
        lugares.append(("Sin datos", 0, set()))

    return lugares


def obtener_etiqueta_similitud(coincidencias):
    if coincidencias == 1:
        return "R"
    return f"R{coincidencias}"


def analizar_archivo(ruta_entrada):
    libro = load_workbook(ruta_entrada, data_only=True)
    hoja = libro.active

    resultados = {}
    for indice_columna in COLUMNAS_A_ANALIZAR:
        letra_columna = get_column_letter(indice_columna)
        contador = Counter()

        for fila in range(1, hoja.max_row + 1):
            numero = normalizar_numero(hoja.cell(row=fila, column=indice_columna).value)
            if numero is not None:
                contador[numero] += 1

        resultados[letra_columna] = contador

    concursos = []
    for fila in range(1, hoja.max_row + 1):
        numeros_fila = {}

        for indice_columna in COLUMNAS_A_ANALIZAR:
            numero = normalizar_numero(hoja.cell(row=fila, column=indice_columna).value)
            if numero is not None:
                numeros_fila[get_column_letter(indice_columna)] = numero

        if numeros_fila:
            numero_concurso = hoja.cell(row=fila, column=COLUMNA_CONCURSO).value
            concursos.append(
                {
                    "fila": fila,
                    "concurso": numero_concurso if numero_concurso is not None else fila,
                    "numeros": numeros_fila,
                }
            )

    return resultados, hoja.title, hoja.max_row, concursos


def crear_hoja_similitudes(libro_reporte, resultados, concursos):
    hoja_similitudes = libro_reporte.create_sheet("Similitudes por concurso")
    hoja_similitudes.append(
        [
            "Lugar analizado",
            "Numero de concurso",
            "Fila origen",
            "Similitud",
            "Coincidencias",
            "Columnas coincidentes",
            "Numeros coincidentes",
        ]
    )

    nombres_lugares = ["Primer lugar", "Segundo lugar", "Tercer lugar"]
    criterios_por_lugar = []

    for indice_lugar, nombre_lugar in enumerate(nombres_lugares):
        criterios = {}
        for columna, contador in resultados.items():
            lugares = obtener_lugares_frecuencia(contador)
            criterios[columna] = lugares[indice_lugar][2]
        criterios_por_lugar.append((nombre_lugar, criterios))

    for nombre_lugar, criterios in criterios_por_lugar:
        for concurso in concursos:
            columnas_coincidentes = []
            numeros_coincidentes = []

            for columna, numeros_frecuentes in criterios.items():
                numero_concurso = concurso["numeros"].get(columna)
                if numero_concurso in numeros_frecuentes:
                    columnas_coincidentes.append(columna)
                    numeros_coincidentes.append(f"{columna}={numero_concurso}")

            coincidencias = len(columnas_coincidentes)
            if coincidencias == 0:
                continue

            hoja_similitudes.append(
                [
                    nombre_lugar,
                    concurso["concurso"],
                    concurso["fila"],
                    obtener_etiqueta_similitud(coincidencias),
                    coincidencias,
                    ", ".join(columnas_coincidentes),
                    ", ".join(numeros_coincidentes),
                ]
            )

    aplicar_estilo_encabezado(hoja_similitudes)
    ajustar_columnas(hoja_similitudes)
    hoja_similitudes.freeze_panes = "A2"


def crear_hoja_tres_lugares(libro_reporte, resultados, concursos):
    hoja_tres_lugares = libro_reporte.create_sheet("Tres lugares juntos")
    hoja_tres_lugares.append(
        [
            "Numero de concurso",
            "Fila origen",
            "Primer lugar encontrado",
            "Segundo lugar encontrado",
            "Tercer lugar encontrado",
            "Total coincidencias",
            "Similitud total",
        ]
    )

    criterios_por_lugar = []
    for indice_lugar in range(3):
        criterios = {}
        for columna, contador in resultados.items():
            lugares = obtener_lugares_frecuencia(contador)
            criterios[columna] = lugares[indice_lugar][2]
        criterios_por_lugar.append(criterios)

    for concurso in concursos:
        coincidencias_por_lugar = []

        for criterios in criterios_por_lugar:
            numeros_coincidentes = []

            for columna, numeros_frecuentes in criterios.items():
                numero_concurso = concurso["numeros"].get(columna)
                if numero_concurso in numeros_frecuentes:
                    numeros_coincidentes.append(f"{columna}={numero_concurso}")

            coincidencias_por_lugar.append(numeros_coincidentes)

        if not all(coincidencias_por_lugar):
            continue

        total_coincidencias = sum(
            len(coincidencias) for coincidencias in coincidencias_por_lugar
        )

        hoja_tres_lugares.append(
            [
                concurso["concurso"],
                concurso["fila"],
                ", ".join(coincidencias_por_lugar[0]),
                ", ".join(coincidencias_por_lugar[1]),
                ", ".join(coincidencias_por_lugar[2]),
                total_coincidencias,
                obtener_etiqueta_similitud(total_coincidencias),
            ]
        )

    aplicar_estilo_encabezado(hoja_tres_lugares)
    ajustar_columnas(hoja_tres_lugares)
    hoja_tres_lugares.freeze_panes = "A2"


def crear_hoja_coincidencias_r7(libro_reporte, resultados, concursos):
    hoja_r7 = libro_reporte.create_sheet("Coincidencias R7")
    hoja_r7.append(
        [
            "Tipo de coincidencia",
            "Numero de concurso",
            "Fila origen",
            "Similitud",
            "Columnas coincidentes",
            "Numeros coincidentes",
        ]
    )

    nombres_lugares = ["Primer lugar", "Segundo lugar", "Tercer lugar"]

    for indice_lugar, nombre_lugar in enumerate(nombres_lugares):
        criterios = {}

        for columna, contador in resultados.items():
            lugares = obtener_lugares_frecuencia(contador)
            criterios[columna] = lugares[indice_lugar][2]

        for concurso in concursos:
            columnas_coincidentes = []
            numeros_coincidentes = []

            for columna, numeros_frecuentes in criterios.items():
                numero_concurso = concurso["numeros"].get(columna)
                if numero_concurso in numeros_frecuentes:
                    columnas_coincidentes.append(columna)
                    numeros_coincidentes.append(f"{columna}={numero_concurso}")

            if len(columnas_coincidentes) != len(list(COLUMNAS_A_ANALIZAR)):
                continue

            hoja_r7.append(
                [
                    nombre_lugar,
                    concurso["concurso"],
                    concurso["fila"],
                    "R7",
                    ", ".join(columnas_coincidentes),
                    ", ".join(numeros_coincidentes),
                ]
            )

    aplicar_estilo_encabezado(hoja_r7)
    ajustar_columnas(hoja_r7)
    hoja_r7.freeze_panes = "A2"


def crear_reporte(ruta_entrada, resultados, nombre_hoja_origen, total_filas, concursos):
    libro_reporte = Workbook()
    hoja_resumen = libro_reporte.active
    hoja_resumen.title = "Resumen"

    hoja_resumen.append(
        [
            "Archivo analizado",
            Path(ruta_entrada).name,
        ]
    )
    hoja_resumen.append(["Hoja analizada", nombre_hoja_origen])
    hoja_resumen.append(["Filas revisadas", total_filas])
    hoja_resumen.append(["Fecha del reporte", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    hoja_resumen.append([])
    hoja_resumen.append(
        [
            "Columna",
            "Primer lugar",
            "Frecuencia 1",
            "Segundo lugar",
            "Frecuencia 2",
            "Tercer lugar",
            "Frecuencia 3",
            "Total de numeros",
            "Numeros diferentes",
        ]
    )

    fila_encabezado_resumen = 6
    for columna, contador in resultados.items():
        total_numeros = sum(contador.values())
        numeros_diferentes = len(contador)

        primer_lugar, segundo_lugar, tercer_lugar = obtener_lugares_frecuencia(contador)

        hoja_resumen.append(
            [
                columna,
                primer_lugar[0],
                primer_lugar[1],
                segundo_lugar[0],
                segundo_lugar[1],
                tercer_lugar[0],
                tercer_lugar[1],
                total_numeros,
                numeros_diferentes,
            ]
        )

    hoja_detalle = libro_reporte.create_sheet("Conteo completo")
    hoja_detalle.append(["Columna", "Numero", "Veces repetido"])

    for columna, contador in resultados.items():
        for numero, frecuencia in sorted(
            contador.items(), key=lambda item: (-item[1], item[0])
        ):
            hoja_detalle.append([columna, numero, frecuencia])

    crear_hoja_similitudes(libro_reporte, resultados, concursos)
    crear_hoja_tres_lugares(libro_reporte, resultados, concursos)
    crear_hoja_coincidencias_r7(libro_reporte, resultados, concursos)

    aplicar_estilo_encabezado(hoja_detalle)
    for celda in hoja_resumen[fila_encabezado_resumen]:
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center")

    ajustar_columnas(hoja_resumen)
    ajustar_columnas(hoja_detalle)
    hoja_resumen.freeze_panes = "A7"
    hoja_detalle.freeze_panes = "A2"

    carpeta_salida = obtener_carpeta_salida()
    carpeta_salida.mkdir(parents=True, exist_ok=True)

    nombre_salida = f"{Path(ruta_entrada).stem}_analisis_frecuencias.xlsx"
    ruta_salida = carpeta_salida / nombre_salida
    libro_reporte.save(ruta_salida)
    return ruta_salida


def seleccionar_archivo():
    ventana = Tk()
    ventana.withdraw()
    ventana.attributes("-topmost", True)

    ruta_archivo = filedialog.askopenfilename(
        title="Selecciona el archivo Excel a analizar",
        filetypes=[
            ("Archivos Excel", "*.xlsx *.xlsm"),
            ("Todos los archivos", "*.*"),
        ],
    )

    ventana.destroy()
    return ruta_archivo


def main():
    ruta_entrada = seleccionar_archivo()
    if not ruta_entrada:
        print("No se selecciono ningun archivo.")
        return

    try:
        resultados, nombre_hoja_origen, total_filas, concursos = analizar_archivo(ruta_entrada)
        ruta_salida = crear_reporte(
            ruta_entrada, resultados, nombre_hoja_origen, total_filas, concursos
        )
    except Exception as error:
        messagebox.showerror("Error", f"No se pudo generar el reporte:\n{error}")
        raise

    mensaje = f"Reporte generado correctamente:\n{ruta_salida}"
    print(mensaje)
    messagebox.showinfo("Analisis completado", mensaje)


if __name__ == "__main__":
    main()
